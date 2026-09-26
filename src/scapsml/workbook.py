"""
workbook -- read the Excel project file into a validated Project.

The workbook is the ONLY file a user edits. Sheets:

  project      key / value settings: base .def, absorber layer, directories
  anchors      one row per anchor composition, one column per material
               property (Eg, chi or IP, eps, Nc, Nv, mun, mup, abs_file)
  sweep        one row per swept parameter: name, layer, mode, values
  layers       optional fixed overrides for any layer in the stack
  experiment   optional measurements to validate against
  reference    generated: every valid parameter name, unit and location

Every problem is reported with the sheet, row and column it came from, and
all problems are collected before failing so a user fixes them in one pass
rather than one per run.
"""

from __future__ import annotations

import math
import re
from dataclasses import dataclass, field
from pathlib import Path

from .registry import DEF, POST, SCRIPT, lookup

__all__ = ["Project", "SweepAxis", "load", "parse_values", "WorkbookError"]

REQUIRED_SHEETS = ("project", "anchors", "sweep")
MODES = ("absolute", "offset", "factor")


class WorkbookError(Exception):
    """One or more problems in the workbook, all listed."""

    def __init__(self, problems):
        self.problems = problems
        super().__init__("\n".join(["workbook has problems:"] +
                                   [f"  - {p}" for p in problems]))


# --------------------------------------------------------------------------
# value parsing
# --------------------------------------------------------------------------
def parse_values(spec):
    """
    Turn a spreadsheet cell into a list of floats.

        1e16                       -> [1e16]
        1e15, 1e16, 1e17           -> explicit list
        20:36:1                    -> 20, 21, ... 36 (inclusive, by step)
        lin:1.0:2.0:5              -> 5 evenly spaced points
        log:1e14:1e18:5            -> 5 points evenly spaced in log10
    """
    if spec is None or (isinstance(spec, str) and not spec.strip()):
        raise ValueError("empty")
    if isinstance(spec, (int, float)):
        return [float(spec)]
    s = str(spec).strip()
    low = s.lower()
    if low.startswith(("lin:", "log:")):
        kind, a, b, n = low.split(":")
        a, b, n = float(a), float(b), int(float(n))
        if n < 1:
            raise ValueError("point count must be >= 1")
        if n == 1:
            return [a]
        if kind == "lin":
            return [a + (b - a) * i / (n - 1) for i in range(n)]
        if a <= 0 or b <= 0:
            raise ValueError("log range needs positive endpoints")
        la, lb = math.log10(a), math.log10(b)
        return [10 ** (la + (lb - la) * i / (n - 1)) for i in range(n)]
    if s.count(":") == 2 and "," not in s:
        a, b, step = (float(x) for x in s.split(":"))
        if step <= 0:
            raise ValueError("step must be positive")
        n = int(math.floor((b - a) / step + 1e-9)) + 1
        return [round(a + i * step, 10) for i in range(n)]
    out = [float(x) for x in re.split(r"[,;\s]+", s) if x.strip()]
    if not out:
        raise ValueError(f"no numbers in {s!r}")
    return out


# --------------------------------------------------------------------------
# data model
# --------------------------------------------------------------------------
@dataclass
class SweepAxis:
    name: str                    # registry name, or "composition"
    values: list
    layer: str | None = None     # None -> absorber
    mode: str = "absolute"       # absolute | offset | factor
    note: str = ""
    row: int = 0

    @property
    def param(self):
        return None if self.name == "composition" else lookup(self.name)

    @property
    def target(self):
        return "composition" if self.name == "composition" else self.param.target


@dataclass
class Project:
    path: Path
    settings: dict
    anchors: list                 # [{"composition": s, "Eg": .., ...}, ...]
    anchor_columns: list
    axes: list                    # [SweepAxis]
    layer_overrides: list         # [(layer, param_name, value)]
    experiment: list              # [{"composition": s, "Eg": .., ...}]
    interpolation: dict = field(default_factory=dict)

    # ---- convenience ---------------------------------------------------
    def setting(self, key, default=None):
        return self.settings.get(key, default)

    @property
    def name(self):
        return self.setting("name", self.path.stem)

    @property
    def absorber(self):
        return self.setting("absorber")

    def resolve(self, key, default=None):
        """A path setting, resolved relative to the workbook's folder."""
        v = self.setting(key, default)
        if v is None:
            return None
        p = Path(str(v))
        return p if p.is_absolute() else (self.path.parent / p).resolve()

    @property
    def base_def(self):
        return self.resolve("base_def")

    def axes_by(self, target):
        return [a for a in self.axes if a.target == target]

    @property
    def composition_axis(self):
        c = [a for a in self.axes if a.name == "composition"]
        return c[0] if c else None


# --------------------------------------------------------------------------
# reading
# --------------------------------------------------------------------------
def _rows(ws):
    """Yield (row_number, [cell values]) skipping blank and comment rows."""
    for r in ws.iter_rows(values_only=False):
        vals = [c.value for c in r]
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in vals):
            continue
        first = vals[0]
        if isinstance(first, str) and first.strip().startswith("#"):
            continue
        yield r[0].row, vals


def _header(ws):
    for n, vals in _rows(ws):
        return n, [str(v).strip() if v is not None else "" for v in vals]
    return None, []


def load(path) -> Project:
    try:
        import openpyxl
    except ImportError as exc:
        raise SystemExit("openpyxl is required:  pip install openpyxl") from exc

    path = Path(path).resolve()
    if not path.is_file():
        raise WorkbookError([f"workbook not found: {path}"])
    wb = openpyxl.load_workbook(path, data_only=True)
    names = {n.lower(): n for n in wb.sheetnames}
    problems = []
    for s in REQUIRED_SHEETS:
        if s not in names:
            problems.append(f"missing required sheet {s!r} "
                            f"(found: {', '.join(wb.sheetnames)})")
    if problems:
        raise WorkbookError(problems)

    # ---- project -------------------------------------------------------
    settings = {}
    for n, vals in _rows(wb[names["project"]]):
        if vals[0] is None:
            continue
        key = str(vals[0]).strip().lower()
        if key in ("key", "setting", "parameter"):
            continue
        settings[key] = vals[1] if len(vals) > 1 else None

    for k in ("base_def", "absorber"):
        if not settings.get(k):
            problems.append(f"project!{k}: required")

    # ---- anchors -------------------------------------------------------
    ws = wb[names["anchors"]]
    hrow, hdr = _header(ws)
    anchors, cols = [], []
    if not hdr or hdr[0].lower() not in ("composition", "s", "x", "s%"):
        problems.append("anchors: first column must be 'composition'")
    else:
        for j, h in enumerate(hdr[1:], start=1):
            if not h:
                continue
            key = h.split("(")[0].strip()
            if key.lower() in ("ip", "abs_file", "abs_onset", "note", "source"):
                cols.append((j, {"ip": "IP"}.get(key.lower(), key.lower())))
                continue
            try:
                p = lookup(key)
            except KeyError as e:
                problems.append(f"anchors!{_col(j)}{hrow} {h!r}: {e}")
                continue
            if p.target != DEF:
                problems.append(f"anchors!{_col(j)}{hrow} {h!r}: {p.name} is a "
                                f"{p.target} parameter; put it in 'sweep' "
                                f"instead -- anchors hold material properties")
                continue
            cols.append((j, p.name))
        for n, vals in _rows(ws):
            if n == hrow:
                continue
            try:
                comp = float(vals[0])
            except (TypeError, ValueError):
                problems.append(f"anchors!A{n}: composition {vals[0]!r} "
                                f"is not a number")
                continue
            rec = {"composition": comp}
            for j, key in cols:
                v = vals[j] if j < len(vals) else None
                if v is None or (isinstance(v, str) and not v.strip()):
                    continue
                if key in ("abs_file", "note", "source"):
                    rec[key] = str(v).strip()
                    continue
                try:
                    rec[key] = float(v)
                except (TypeError, ValueError):
                    problems.append(f"anchors!{_col(j)}{n}: {v!r} is not a number")
            anchors.append(rec)
        anchors.sort(key=lambda r: r["composition"])
        comps = [a["composition"] for a in anchors]
        if len(set(comps)) != len(comps):
            problems.append("anchors: duplicate composition values")
        if not anchors:
            problems.append("anchors: at least one row is required")
        colnames = [k for _, k in cols]
        if "chi" in colnames and "IP" in colnames:
            problems.append("anchors: give either 'chi' or 'IP', not both -- "
                            "chi is derived as IP - Eg when IP is given")
        if "Eg" not in colnames:
            problems.append("anchors: an 'Eg' column is required")

    # ---- interpolation (optional) --------------------------------------
    interp = {}
    if "interpolation" in names:
        for n, vals in _rows(wb[names["interpolation"]]):
            if vals[0] is None or str(vals[0]).lower() in ("property", "parameter"):
                continue
            m = str(vals[1] or "").strip().lower()
            if m not in ("linear", "log", "pchip"):
                problems.append(f"interpolation!B{n}: method must be linear, "
                                f"log or pchip, got {vals[1]!r}")
                continue
            interp[str(vals[0]).strip()] = m

    # ---- sweep ---------------------------------------------------------
    axes = []
    ws = wb[names["sweep"]]
    hrow, hdr = _header(ws)
    idx = {h.lower(): j for j, h in enumerate(hdr)}
    for need in ("parameter", "values"):
        if need not in idx:
            problems.append(f"sweep: missing column {need!r} "
                            f"(have: {', '.join(h for h in hdr if h)})")
    if "parameter" in idx and "values" in idx:
        seen = set()
        for n, vals in _rows(ws):
            if n == hrow:
                continue
            g = lambda k: vals[idx[k]] if k in idx and idx[k] < len(vals) else None
            name = str(g("parameter") or "").strip()
            if not name:
                continue
            canon = "composition" if name.lower() in (
                "composition", "s", "x", "s%") else None
            if canon is None:
                try:
                    canon = lookup(name).name
                except KeyError as e:
                    problems.append(f"sweep!A{n} {name!r}: {e}")
                    continue
            try:
                values = parse_values(g("values"))
            except ValueError as e:
                problems.append(f"sweep!{_col(idx['values'])}{n} "
                                f"({name}): cannot read values -- {e}")
                continue
            mode = str(g("mode") or "absolute").strip().lower()
            if mode not in MODES:
                problems.append(f"sweep!{_col(idx.get('mode', 0))}{n}: mode "
                                f"must be one of {MODES}, got {mode!r}")
                continue
            if mode != "absolute" and canon == "composition":
                problems.append(f"sweep!A{n}: composition must be absolute")
                continue
            if mode != "absolute" and lookup(canon).target != DEF:
                problems.append(f"sweep!A{n} {canon}: offset/factor modes only "
                                f"apply to material properties that come from "
                                f"the anchors")
                continue
            layer = g("layer")
            layer = str(layer).strip() if layer not in (None, "") else None
            key = (canon, layer)
            if key in seen:
                problems.append(f"sweep!A{n}: {canon} swept twice"
                                + (f" for layer {layer}" if layer else ""))
                continue
            seen.add(key)
            axes.append(SweepAxis(canon, values, layer, mode,
                                  str(g("note") or ""), n))
        if not any(a.name == "composition" for a in axes):
            problems.append("sweep: no 'composition' row -- say which "
                            "compositions to simulate")

    # ---- layers (optional) ---------------------------------------------
    overrides = []
    if "layers" in names:
        for n, vals in _rows(wb[names["layers"]]):
            if vals[0] is None or str(vals[0]).lower() == "layer":
                continue
            layer, pname, value = (vals + [None, None, None])[:3]
            if pname is None:
                continue
            if str(pname).strip().lower() == "absorption":
                overrides.append((str(layer).strip(), "absorption",
                                  str(value).strip() if value else "model"))
                continue
            try:
                p = lookup(pname)
                overrides.append((str(layer).strip(), p.name, float(value)))
            except KeyError as e:
                problems.append(f"layers!B{n} {pname!r}: {e}")
            except (TypeError, ValueError):
                problems.append(f"layers!C{n}: {value!r} is not a number")

    # ---- experiment (optional) -----------------------------------------
    experiment = []
    if "experiment" in names:
        ws = wb[names["experiment"]]
        hrow, hdr = _header(ws)
        if hdr:
            for n, vals in _rows(ws):
                if n == hrow:
                    continue
                rec = {}
                for j, h in enumerate(hdr):
                    if not h or j >= len(vals) or vals[j] is None:
                        continue
                    key = h.split("(")[0].strip()
                    if key.lower() in ("source", "note", "reference"):
                        rec[key.lower()] = str(vals[j])
                        continue
                    try:
                        rec["composition" if j == 0 else key] = float(vals[j])
                    except (TypeError, ValueError):
                        problems.append(f"experiment!{_col(j)}{n}: "
                                        f"{vals[j]!r} is not a number")
                if rec:
                    experiment.append(rec)

    # ---- cross-checks against the base .def ----------------------------
    proj = Project(path, settings, anchors, [k for _, k in cols], axes,
                   overrides, experiment, interp)
    if proj.base_def and not proj.base_def.is_file():
        problems.append(f"project!base_def: file not found: {proj.base_def}")
    elif proj.base_def:
        from .defio import DefFile, DefError
        d = DefFile.load(proj.base_def)
        if proj.absorber and proj.absorber not in d.layers:
            problems.append(f"project!absorber: {proj.absorber!r} is not a layer "
                            f"in {proj.base_def.name}. Layers: {d.layers}")
        for a in axes:
            if a.layer and a.layer not in d.layers and "/" not in a.layer:
                problems.append(f"sweep!row {a.row} ({a.name}): layer "
                                f"{a.layer!r} not in the .def. Layers: {d.layers}")
        for layer, pname, _ in overrides:
            if layer not in d.layers:
                problems.append(f"layers: {layer!r} is not a layer in the .def. "
                                f"Layers: {d.layers}")

    if problems:
        raise WorkbookError(problems)
    return proj


def _col(j):
    s = ""
    j += 1
    while j:
        j, r = divmod(j - 1, 26)
        s = chr(65 + r) + s
    return s
